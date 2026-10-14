"""Keep `G:\\My Drive\\school-out` organised so the Gemini app has the WHOLE course to read.

(2026-10-12) "organize the drive so gemini has proper context — I used nebula the whole day
today and there were times that it had incomplete context to accomplish tasks."

Two causes, two fixes:
  1. Blackboard files reached Drive only when a session downloaded one by hand. `lms.sh mirror`
     (bb_mirror.py) now copies every course's files + page text; this module publishes them to
     `<CODE>\\materials\\blackboard\\<Blackboard folder>\\…`.
  2. Nothing told Gemini what was where. Each course gets `<CODE> course guide.docx` (what is due,
     what the instructor said, every file by folder) and the root gets `00 school-out guide.docx`.
     The ✨ prompts name the guide, so Gemini opens it first.

Layout (the student's own files were sorted into it on 2026-10-12 — vault school/context/drive-reorg-2026-10-12.md):
    <CODE>\\<CODE> course guide.docx · materials\\blackboard\\ · materials\\ · my-work\\ · study\\

Copying goes WSL → C:\\Temp\\nebula\\stage → robocopy into Drive for Desktop's G:, the same
path studypack.publish uses. robocopy runs WITHOUT /MIR or /PURGE: nothing on Drive is deleted.
School data only — the same privacy rule as the judge (personal mail never reaches a guide).
"""
import json
import re
import shutil
import subprocess
from datetime import timedelta
from pathlib import Path

from django.conf import settings
from django.utils import timezone

from .judge import COURSES

POWERSHELL = "/mnt/c/Windows/System32/WindowsPowerShell/v1.0/powershell.exe"
STAGE = Path("/mnt/c/Temp/nebula/stage")
STAGE_WIN = "C:\\Temp\\nebula\\stage"
GUIDE = "{code} course guide.docx"
ROOT_GUIDE = "00 school-out guide.docx"
SKIP_LIST = {"desktop.ini", "Thumbs.db"}


def mirror_dir():
    return settings.DATA_DIR / "materials"


def schedule_row(code):
    """(name, units, times, rooms, instructor) from school/schedule.md's course table."""
    f = settings.SCHOOL_DIR / "schedule.md"
    if not f.exists():
        return None
    for line in f.read_text().splitlines():
        cells = [c.strip().strip("`*") for c in line.strip().strip("|").split("|")]
        if len(cells) >= 7 and cells[1] == code:
            return {"name": cells[2], "units": cells[3], "times": cells[4], "rooms": cells[5],
                    "instructor": cells[6].replace("⚠", "").strip()}
    return None


def listing_cache():
    return settings.DATA_DIR / "drive_listing.json"


def drive_listing():
    """{CODE: [relative paths]} of what is on Drive now (one PowerShell call).

    Also caches {CODE: [[path, modified-UTC]]} for the ✨ button, which cannot afford a 30-second
    PowerShell walk per tap (2026-10-13: it never named the reviewers sitting in study\\)."""
    root = settings.DRIVE_OUT
    ps = (f"Get-ChildItem -Recurse -File -LiteralPath '{root}' | ForEach-Object {{ "
          f"$_.FullName.Substring({len(root) + 1}) + '|' + "
          f"$_.LastWriteTimeUtc.ToString('yyyy-MM-ddTHH:mm:ssZ') }}")
    p = subprocess.run([POWERSHELL, "-NoProfile", "-Command", ps], capture_output=True, text=True,
                       timeout=180)
    out, stamped = {}, {}
    for line in p.stdout.splitlines():
        rel, _, when = line.strip().rpartition("|")
        if not rel:                              # no stamp on the line: keep it, undated
            rel, when = line.strip(), ""
        parts = rel.split("\\")
        if len(parts) >= 2 and parts[0] in COURSES and parts[-1] not in SKIP_LIST:
            out.setdefault(parts[0], []).append("\\".join(parts[1:]))
            stamped.setdefault(parts[0], []).append(["\\".join(parts[1:]), when])
    if p.returncode == 0 and stamped:            # a failed walk must not blank a good cache
        listing_cache().write_text(json.dumps(stamped))
    return out


def _manifest(code):
    f = mirror_dir() / code / "_manifest.json"
    return json.loads(f.read_text()) if f.exists() else {}


def _chat_manifest(code):
    f = mirror_dir() / code / "_chat_manifest.json"
    return json.loads(f.read_text()) if f.exists() else {}


def _when(dt):
    return timezone.localtime(dt).strftime("%a %d %b, %I:%M %p").replace(" 0", " ") if dt else "no date"


def _section_files(listing, prefix):
    return sorted(p for p in listing if p.lower().startswith(prefix.lower() + "\\"))


def guide_lines(code, listing=None):
    """The guide as (style, text) pairs — kept separate from docx so it can be tested."""
    from .done import done_keys
    from .models import CalEvent, Item, Message
    from .present import AGE_RULE, age, clean_title, today_line, upcoming
    now = timezone.now()
    name = COURSES.get(code, code)
    row = schedule_row(code) or {}
    marked = done_keys()
    L = [("title", f"{code} — {row.get('name') or name}"),
         ("p", f"Course guide for Gemini · built by Nebula {_when(now)}. {today_line()} Every fact "
               f"below comes from Blackboard, Classroom, Chat, the student's calendar or the files in this folder.")]
    if row:
        L += [("h", "The class"),
              ("li", f"Instructor: {row['instructor']}"), ("li", f"Meets: {row['times']}"),
              ("li", f"Rooms: {row['rooms']} (VC01 = online)"), ("li", f"Units: {row['units']}")]
    L += [("h", "How this folder is organised"),
          ("li", "materials\\blackboard\\ — every file and page the instructor posted on Blackboard, "
                 "in Blackboard's own folders. A .txt file is the text of a Blackboard page or the "
                 "instructions of an assignment."),
          ("li", "materials\\chat\\ — files the instructor posted in the course's Google Chat space."),
          ("li", "materials\\classroom\\ — files from Google Classroom posts that are not also on Blackboard."),
          ("li", "materials\\ — instructor files saved by hand (handouts, rubrics, book chapters)."),
          ("li", "my-work\\ — what the student made and submitted."),
          ("li", "study\\ — reviewers, cram sheets, build cards and the study pack."),
          ("p", "When asked about a task: find it under \"What is due\", then open the matching "
                "file in materials\\blackboard\\ before answering. Quote the instructor's file.")]
    items = [i for i in Item.objects.filter(course=code, submitted=False, due__gte=now - timedelta(days=7))
             if i.uid not in marked]
    pins = list(CalEvent.objects.filter(start__gte=now - timedelta(days=1), start__lt=now + timedelta(days=30)))
    L.append(("h", "What is due (not yet submitted)"))
    ups = upcoming(code, now - timedelta(days=7), items, pins)
    if ups:
        for due, title in ups:
            past = " — OVERDUE" if due < now else ""
            L.append(("li", f"{_when(due)} · {title}{past}"))
    else:
        L.append(("li", "Nothing dated is open for this course."))
    msgs = Message.objects.filter(course=code, gemini_eligible=True, when__gte=now - timedelta(days=21)
                                  ).filter(judge_relevant=True).order_by("-when")[:12]
    L.append(("h", "What the instructor said recently (last 3 weeks)"))
    L.append(("p", AGE_RULE))
    for m in msgs:
        body = " ".join((m.body or "").split())[:900]
        L.append(("li", f"{timezone.localtime(m.when):%a %d %b} ({age(m.when)}) · {m.source_long}"
                        f"{' · ' + m.author if m.author else ''}: {body}"))
    if not msgs:
        L.append(("li", "No flagged announcements in the last 3 weeks."))
    man = _manifest(code)
    files = man.get("files", [])
    L.append(("h", f"Blackboard files ({len(files)}) — materials\\blackboard\\"))
    for f in sorted(files, key=lambda f: (f["path"].lower())):
        posted = (f.get("created") or "")[:10]
        L.append(("li", f"{f['path'].replace('/', chr(92))}  —  \"{clean_title(f['title'])}\""
                        f"{', posted ' + posted if posted else ''}"))
    if not man:
        L.append(("li", "The Blackboard mirror has not run for this course yet."))
    vids = man.get("skipped", [])
    if vids:
        L.append(("h", "On Blackboard only (videos and files over 100 MB — not copied to Drive)"))
        L += [("li", f"{v.get('folder') or '(top level)'} / {v['file']}") for v in vids]
    cm = _chat_manifest(code)
    if cm.get("files") or cm.get("drive_only"):
        L.append(("h", f"Files the instructor posted in Google Chat ({len(cm.get('files', []))}) — materials\\chat\\"))
        for f in cm.get("files", []):
            L.append(("li", f"{f['path'].split('/', 1)[-1]}  —  {f.get('author', '')}, {f.get('posted', '')}"
                            + (f": \"{f['message'][:140]}\"" if f.get("message") else "")))
        for f in cm.get("drive_only", []):
            L.append(("li", f"(Drive link, open in Chat) {f['file']}  —  {f.get('posted', '')}"))
    links = man.get("links", [])
    if links:
        L.append(("h", "Links the instructor posted"))
        L += [("li", f"{l['title']}: {l['url']}") for l in links]
    listing = listing or []
    for folder, label in (("materials\\classroom", "Google Classroom files — materials\\classroom\\"),
                          ("materials", "Saved instructor files — materials\\"),
                          ("my-work", "the student's work — my-work\\"), ("study", "Study files — study\\")):
        got = [p for p in _section_files(listing, folder)
               if not p.lower().startswith(("materials\\blackboard\\", "materials\\chat\\"))
               and (folder != "materials" or not p.lower().startswith("materials\\classroom\\"))]
        L.append(("h", f"{label} ({len(got)})"))
        L += [("li", p) for p in got] or [("li", "— empty —")]
    return L


def write_docx(lines, path):
    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    for style, text in lines:
        if style == "title":
            doc.add_heading(text, level=0)
        elif style == "h":
            doc.add_heading(text, level=2)
        elif style == "li":
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def root_lines():
    L = [("title", "school-out — the student's school Drive"),
         ("p", f"Organised by Nebula · {_when(timezone.now())}. One folder per course."),
         ("h", "Start here"),
         ("p", "Each course folder has \"<CODE> course guide.docx\": what is due, what the instructor "
               "said, and every file in that course by folder. Open the guide first."),
         ("h", "Courses")]
    for code, name in COURSES.items():
        row = schedule_row(code) or {}
        L.append(("li", f"{code} — {row.get('name') or name}"
                        f"{' · ' + row['instructor'] if row.get('instructor') else ''}"))
    L += [("h", "Other folders"),
          ("li", "_shared — files for several courses (exam schedule); _shared\\chat holds the BSCS 2627A program-wide Chat posts."),
          ("li", "_general, _unsorted, _other — older notes and loose files, not course material.")]
    return L


def stage_mirror(code):
    """Copy the local Blackboard mirror for CODE into the Windows staging tree (changed files only)."""
    src = mirror_dir() / code
    if not src.exists():
        return 0
    dst = STAGE / code / "materials" / "blackboard"
    n = 0
    for f in src.rglob("*"):
        if f.is_dir() or f.name.startswith(("_manifest", "_chat_manifest")):
            continue
        rel = f.relative_to(src)
        # Instructors' Chat uploads → materials\chat (2026-10-13); _shared → its own top folder.
        if rel.parts[0] == "_chat":
            t = STAGE / code / ("chat" if code == "_shared" else "materials/chat") / Path(*rel.parts[1:])
        else:
            t = dst / rel
        if t.exists() and t.stat().st_size == f.stat().st_size and t.stat().st_mtime >= f.stat().st_mtime:
            continue
        t.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(f, t)
        n += 1
    return n


def robocopy():
    """Stage → Drive. No /MIR, no /PURGE: robocopy only adds and updates."""
    ps = (f"robocopy '{STAGE_WIN}' '{settings.DRIVE_OUT}' /E /XO /R:1 /W:2 /NP /NJH /NDL /NFL; "
          f"exit $LASTEXITCODE")
    p = subprocess.run([POWERSHELL, "-NoProfile", "-Command", ps], capture_output=True, text=True,
                       timeout=1800)
    if p.returncode >= 8:                       # robocopy: 0-7 = success, 8+ = a copy failed
        raise RuntimeError(f"robocopy exit {p.returncode}: {(p.stdout + p.stderr).strip()[-300:]}")
    m = re.search(r"Files\s*:\s*(\d+)\s+(\d+)", p.stdout)
    return int(m.group(2)) if m else 0


def publish(codes=None):
    """Stage the mirror + fresh guides, then robocopy to Drive. Returns a one-line summary."""
    codes = codes or list(COURSES)
    STAGE.mkdir(parents=True, exist_ok=True)
    staged = sum(stage_mirror(c) for c in list(codes) + ["_shared"])
    listing = drive_listing()
    for c in codes:
        write_docx(guide_lines(c, listing.get(c, [])), STAGE / c / GUIDE.format(code=c))
    write_docx(root_lines(), STAGE / ROOT_GUIDE)
    copied = robocopy()
    return f"{staged} mirror file(s) staged · {copied} copied to Drive · {len(codes)} guides"
